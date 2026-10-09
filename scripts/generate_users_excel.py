import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# Sheet 1: Active Users Directory
ws = wb.active
ws.title = 'Users & Login Details'

# Title banner
ws.merge_cells('A1:H1')
title_cell = ws['A1']
title_cell.value = 'RODITTE FLEET MANAGEMENT — OFFICIAL USERS & CREDENTIALS DIRECTORY'
title_cell.font = Font(name='Segoe UI', size=14, bold=True, color='FFFFFF')
title_cell.fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
title_cell.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 40

# Subtitle banner
ws.merge_cells('A2:H2')
sub_cell = ws['A2']
sub_cell.value = 'Confidential • Updated October 2026 • Phone Number Username Standard with Universal Password Reset Permission'
sub_cell.font = Font(name='Segoe UI', size=9, italic=True, color='E0E7FF')
sub_cell.fill = PatternFill(start_color='1D4ED8', end_color='1D4ED8', fill_type='solid')
sub_cell.alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 24

# Header row
headers = [
    'Full Name',
    'Role / Designation',
    'Login Username',
    'Current / Default Password',
    'Password Reset Permission',
    'Assigned Vehicle / Depot',
    'Mobile Number (+91)',
    'Status'
]

ws.row_dimensions[3].height = 28
header_fill = PatternFill(start_color='2563EB', end_color='2563EB', fill_type='solid')
header_font = Font(name='Segoe UI', size=11, bold=True, color='FFFFFF')
border_thin = Border(
    left=Side(style='thin', color='CBD5E1'),
    right=Side(style='thin', color='CBD5E1'),
    top=Side(style='thin', color='CBD5E1'),
    bottom=Side(style='thin', color='CBD5E1')
)

for col_idx, h in enumerate(headers, 1):
    cell = ws.cell(row=3, column=col_idx, value=h)
    cell.fill = header_fill
    cell.font = header_font
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = border_thin

users_data = [
    # Super Admin
    ('Edwin', 'Super Administrator', 'edwin', 'Admin@12345', 'Enabled (In-App Menu)', 'All Depots & Fleets', '+91 98450 00001', 'ACTIVE'),
    # Godown Managers
    ('Rajan Pillai', 'Godown Manager', '9845022001', 'Staff@12345', 'Enabled (In-App Menu)', 'Peenya Central Godown', '+91 98450 22001', 'ACTIVE'),
    ('Manjunath Gowda', 'Godown Manager', '9845022002', 'Staff@12345', 'Enabled (In-App Menu)', 'Whitefield Depot', '+91 98450 22002', 'ACTIVE'),
    # Sales Staff
    ('Ananya Sharma', 'Sales Staff', '9845033001', 'Staff@12345', 'Enabled (In-App Menu)', 'West Division Dispatch', '+91 98450 33001', 'ACTIVE'),
    ('Vikram Joshi', 'Sales Staff', '9845033002', 'Staff@12345', 'Enabled (In-App Menu)', 'East Division Dispatch', '+91 98450 33002', 'ACTIVE'),
    # Drivers
    ('Kiran Kumar', 'Fleet Driver', '9845011001', 'Driver@12345', 'Enabled (Lock Icon)', 'Tata Ace Gold (KA-04-AB-1234)', '+91 98450 11001', 'AVAILABLE'),
    ('Ramesh Babu', 'Fleet Driver', '9845011002', 'Driver@12345', 'Enabled (Lock Icon)', 'Eicher Pro 1049 (KA-05-CD-5678)', '+91 98450 11002', 'BUSY - IN RUN'),
    ('Sunil V', 'Fleet Driver', '9845011003', 'Driver@12345', 'Enabled (Lock Icon)', 'Mahindra Bolero Maxi (KA-51-EF-9012)', '+91 98450 11003', 'AVAILABLE'),
    ('Anand Rao', 'Fleet Driver', '9845011004', 'Driver@12345', 'Enabled (Lock Icon)', 'Piaggio Ape Electric (KA-03-GH-3456)', '+91 98450 11004', 'OFF DUTY'),
    ('Suresh Nair', 'Fleet Driver', '9845011999', 'Driver@12345', 'Enabled (Lock Icon)', 'Unassigned (Driver Choice)', '+91 98450 11999', 'ACTIVE')
]

row_bg_even = PatternFill(start_color='F8FAFC', end_color='F8FAFC', fill_type='solid')
row_bg_odd = PatternFill(start_color='FFFFFF', end_color='FFFFFF', fill_type='solid')
font_normal = Font(name='Segoe UI', size=10, color='1E293B')
font_bold = Font(name='Segoe UI', size=10, bold=True, color='0F172A')
font_username = Font(name='Segoe UI', size=10, bold=True, color='1D4ED8')
font_pwd = Font(name='Segoe UI', size=10, bold=True, color='15803D')
font_busy = Font(name='Segoe UI', size=10, bold=True, color='DC2626')
font_avail = Font(name='Segoe UI', size=10, bold=True, color='16A34A')

for row_idx, u in enumerate(users_data, 4):
    ws.row_dimensions[row_idx].height = 24
    bg = row_bg_even if row_idx % 2 == 0 else row_bg_odd
    
    ws.cell(row=row_idx, column=1, value=u[0]).font = font_bold
    ws.cell(row=row_idx, column=2, value=u[1]).font = font_normal
    ws.cell(row=row_idx, column=3, value=u[2]).font = font_username
    ws.cell(row=row_idx, column=4, value=u[3]).font = font_pwd
    ws.cell(row=row_idx, column=5, value=u[4]).font = font_normal
    ws.cell(row=row_idx, column=6, value=u[5]).font = font_normal
    ws.cell(row=row_idx, column=7, value=u[6]).font = font_normal
    
    status_cell = ws.cell(row=row_idx, column=8, value=u[7])
    if 'BUSY' in u[7]:
        status_cell.font = font_busy
    elif 'AVAIL' in u[7] or 'ACTIVE' in u[7]:
        status_cell.font = font_avail
    else:
        status_cell.font = font_normal

    for c in range(1, 9):
        cell = ws.cell(row=row_idx, column=c)
        cell.fill = bg
        cell.border = border_thin
        if c in (2, 3, 4, 5, 7, 8):
            cell.alignment = Alignment(horizontal='center', vertical='center')
        else:
            cell.alignment = Alignment(horizontal='left', vertical='center')

# Sheet 2: Login Instructions & Rules
ws2 = wb.create_sheet(title='Login Rules & Guide')
ws2.merge_cells('A1:F1')
ws2['A1'].value = 'RODITTE FLEET MANAGEMENT — AUTHENTICATION & OPERATIONS POLICY'
ws2['A1'].font = Font(name='Segoe UI', size=13, bold=True, color='FFFFFF')
ws2['A1'].fill = PatternFill(start_color='1E3A8A', end_color='1E3A8A', fill_type='solid')
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 35

rules = [
    ('Feature / Rule', 'Detailed Description & Implementation'),
    ('Phone Number Usernames', 'All newly created Sales Staff, Godown Managers, and Drivers use their 10-digit mobile number as their login username. Super Admin Edwin can log in with "edwin" or his registered number.'),
    ('Default Passwords', 'Upon onboarding: Staff accounts default to "Staff@12345", Driver accounts default to "Driver@12345", Admin account defaults to "Admin@12345".'),
    ('Universal Password Change', 'Every user has in-app permission to change their password anytime. Admin and Staff access "Change Password" via the 3-dots top menu. Drivers access it via the lock icon on the top bar.'),
    ('Driver Vehicle Selection', 'When Edwin adds a driver, vehicle assignment is OPTIONAL (defaults to "None - Driver Choice"). Drivers choose any available vehicle based on the order and parcel size assigned by Godown Managers.'),
    ('Busy Vehicle Locking', 'Vehicles actively on a trip are marked "BUSY - IN RUN" and disabled from being chosen by other drivers until the current delivery trip is finished.')
]

ws2.row_dimensions[3].height = 26
ws2.cell(row=3, column=1, value=rules[0][0]).fill = header_fill
ws2.cell(row=3, column=1).font = header_font
ws2.cell(row=3, column=1).alignment = Alignment(horizontal='center', vertical='center')

ws2.cell(row=3, column=2, value=rules[0][1]).fill = header_fill
ws2.cell(row=3, column=2).font = header_font
ws2.cell(row=3, column=2).alignment = Alignment(horizontal='center', vertical='center')

for r_idx, (k, v) in enumerate(rules[1:], 4):
    ws2.row_dimensions[r_idx].height = 36
    c1 = ws2.cell(row=r_idx, column=1, value=k)
    c1.font = font_bold
    c1.fill = row_bg_even if r_idx % 2 == 0 else row_bg_odd
    c1.alignment = Alignment(horizontal='left', vertical='center')
    c1.border = border_thin
    
    c2 = ws2.cell(row=r_idx, column=2, value=v)
    c2.font = font_normal
    c2.fill = row_bg_even if r_idx % 2 == 0 else row_bg_odd
    c2.alignment = Alignment(horizontal='left', vertical='center', wrap_text=True)
    c2.border = border_thin

ws2.column_dimensions['A'].width = 30
ws2.column_dimensions['B'].width = 85

# Auto-fit sheet 1 columns
for col in ws.columns:
    max_len = max(len(str(cell.value or '')) for cell in col)
    col_letter = get_column_letter(col[0].column)
    ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

dest1 = 'C:/Users/Sreejith BS/Downloads/Roditte_Fleet_Users_Login_Details.xlsx'
dest2 = 'C:/Users/Sreejith BS/.gemini/antigravity/scratch/fleet-delivery-platform/Roditte_Fleet_Users_Login_Details.xlsx'
wb.save(dest1)
wb.save(dest2)
print('Workbook saved successfully to Downloads and project directory!')
